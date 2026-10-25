"""Build a 16:9 PPTX from the rendered slide PNGs (one full-bleed picture per slide) with speaker notes.

    ../../.venv/Scripts/python build_pptx.py ../deck_render
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

src = Path(sys.argv[1])
prs = Presentation()
prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)          # 13.333 x 7.5 in (16:9)
notes = json.loads((src / "notes.json").read_text(encoding="utf-8"))
for i, png in enumerate(sorted(src.glob("slide_*.png"))):
    s = prs.slides.add_slide(prs.slide_layouts[6])                         # blank
    s.shapes.add_picture(str(png), 0, 0, width=prs.slide_width, height=prs.slide_height)
    if i < len(notes) and notes[i]:
        s.notes_slide.notes_text_frame.text = notes[i]
out = src / "QueueShield_pitch_deck.pptx"
prs.save(out)
print(json.dumps({"pptx": str(out), "slides": len(prs.slides)}))
